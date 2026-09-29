from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os
from datetime import datetime

from app.database.base import get_db
from app.models.deal import Deal
from app.models.search import SearchQuery

router = APIRouter()


@router.get("/excel/{search_id}")
async def export_to_excel(
    search_id: int,
    db: AsyncSession = Depends(get_db)
):
    """Export search results to formatted Excel file"""
    
    # Get search and deals
    result = await db.execute(
        select(SearchQuery).where(SearchQuery.id == search_id)
    )
    search = result.scalar_one_or_none()
    
    if not search:
        raise HTTPException(status_code=404, detail="Search not found")
    
    result = await db.execute(
        select(Deal).where(Deal.search_query_id == search_id).order_by(Deal.price)
    )
    deals = result.scalars().all()
    
    if not deals:
        raise HTTPException(status_code=404, detail="No deals found for this search")
    
    # Create Excel file
    filename = await create_excel_report(search, deals)
    
    return FileResponse(
        filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"travel_deals_{search_id}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )


async def create_excel_report(search: SearchQuery, deals: list) -> str:
    """Create formatted Excel report with multiple sheets"""
    
    filename = f"/tmp/travel_deals_{search.id}.xlsx"
    
    # Convert deals to DataFrames
    hotels_data = []
    flights_data = []
    
    for deal in deals:
        if deal.deal_type == "hotel":
            hotels_data.append({
                'Hotel Name': deal.hotel_name,
                'Price': deal.price,
                'Rating': deal.hotel_rating,
                'Room Type': deal.room_type,
                'Check In': deal.check_in_date,
                'Check Out': deal.check_out_date,
                'Location': deal.location,
                'Source': deal.source,
                'URL': deal.url,
            })
        else:
            flights_data.append({
                'Airline': deal.airline,
                'Price': deal.price,
                'Departure': deal.departure_time,
                'Arrival': deal.arrival_time,
                'Duration': deal.duration,
                'Stops': deal.stops,
                'Date': deal.check_in_date,
                'Destination': deal.location,
                'Source': deal.source,
                'URL': deal.url,
            })
    
    # Create Excel writer
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        
        # Summary sheet
        summary_data = {
            'Search Criteria': [
                'Location',
                'Check In',
                'Check Out',
                'Total Deals',
                'Hotels Found',
                'Flights Found',
                'Lowest Price',
                'Highest Price',
                'Average Price',
            ],
            'Value': [
                search.location,
                search.check_in_date,
                search.check_out_date,
                len(deals),
                len(hotels_data),
                len(flights_data),
                f"${min([d.price for d in deals]):.2f}",
                f"${max([d.price for d in deals]):.2f}",
                f"${sum([d.price for d in deals]) / len(deals):.2f}",
            ]
        }
        df_summary = pd.DataFrame(summary_data)
        df_summary.to_excel(writer, sheet_name='Summary', index=False)
        
        # Hotels sheet
        if hotels_data:
            df_hotels = pd.DataFrame(hotels_data)
            df_hotels.to_excel(writer, sheet_name='Hotels', index=False)
        
        # Flights sheet
        if flights_data:
            df_flights = pd.DataFrame(flights_data)
            df_flights.to_excel(writer, sheet_name='Flights', index=False)
        
        # All Deals sheet
        all_deals_data = []
        for deal in deals:
            all_deals_data.append({
                'Type': deal.deal_type.capitalize(),
                'Name': deal.name,
                'Price': deal.price,
                'Location': deal.location,
                'Date': deal.check_in_date,
                'Source': deal.source,
            })
        df_all = pd.DataFrame(all_deals_data)
        df_all.to_excel(writer, sheet_name='All Deals', index=False)
    
    # Format the workbook
    wb = load_workbook(filename)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Style header row
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Add currency formatting to Price columns
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if ws.cell(1, cell.column).value == 'Price':
                    cell.number_format = '$#,##0.00'
    
    wb.save(filename)
    return filename


@router.get("/csv/{search_id}")
async def export_to_csv(
    search_id: int,
    db: AsyncSession = Depends(get_db)
):
    """Export search results to CSV file"""
    
    result = await db.execute(
        select(Deal).where(Deal.search_query_id == search_id)
    )
    deals = result.scalars().all()
    
    if not deals:
        raise HTTPException(status_code=404, detail="No deals found")
    
    # Create DataFrame
    deals_data = [deal.to_dict() for deal in deals]
    df = pd.DataFrame(deals_data)
    
    # Save to CSV
    filename = f"/tmp/travel_deals_{search_id}.csv"
    df.to_csv(filename, index=False)
    
    return FileResponse(
        filename,
        media_type="text/csv",
        filename=f"travel_deals_{search_id}.csv"
    )
